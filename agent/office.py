#!/usr/bin/env python3
"""تحويل نص Markdown بسيط لملف Word يدعم العربية (من اليمين لليسار).

  pip install python-docx
  python3 agent/office.py result.md out.docx [--ltr]

يدعم: # عناوين، - نقاط، 1. ترقيم، **غامق**، وأسطر عادية.
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ARABIC = re.compile(r"[؀-ۿ]")


def _bidi(paragraph, rtl):
    ppr = paragraph._p.get_or_add_pPr()
    b = OxmlElement("w:bidi")
    b.set(qn("w:val"), "1" if rtl else "0")
    ppr.append(b)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT


def _runs(paragraph, text, rtl, size=None):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = i % 2 == 1
        rpr = run._r.get_or_add_rPr()
        if rtl or ARABIC.search(part):
            rtl_el = OxmlElement("w:rtl")
            rpr.append(rtl_el)
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.append(fonts)
            fonts.set(qn("w:cs"), "Arial")
            szcs = OxmlElement("w:szCs")
            szcs.set(qn("w:val"), str(int((size or 11) * 2)))
            rpr.append(szcs)
        if size:
            run.font.size = Pt(size)


def md_to_docx(md, out, rtl=True):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    for line in md.splitlines():
        s = line.rstrip()
        if not s.strip():
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            level = len(m.group(1))
            p = doc.add_heading(level=level)
            _bidi(p, rtl)
            _runs(p, m.group(2), rtl, size={1: 18, 2: 14, 3: 12}[level])
            continue
        m = re.match(r"^\s*[-•*]\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _bidi(p, rtl)
            _runs(p, m.group(1), rtl)
            continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            _bidi(p, rtl)
            _runs(p, m.group(1), rtl)
            continue
        p = doc.add_paragraph()
        _bidi(p, rtl)
        _runs(p, s, rtl)
    doc.save(out)
    return out


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as fh:
        md_to_docx(fh.read(), sys.argv[2], rtl="--ltr" not in sys.argv)
    print("saved", sys.argv[2])
